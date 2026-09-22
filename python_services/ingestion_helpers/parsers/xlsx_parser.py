from pathlib import Path
from openpyxl import load_workbook


def extract_xlsx_text(file_path: str) -> list[dict]:
    """
    Extract text from each row of every sheet in an XLSX file.
    """

    xlsx_path = Path(file_path)

    if not xlsx_path.exists():
        raise FileNotFoundError(f"XLSX file not found: {file_path}")

    workbook = load_workbook(
        filename=xlsx_path,
        read_only=True,
        data_only=True
    )

    rows = []

    for sheet in workbook.worksheets:

        for row_number, row in enumerate(
            sheet.iter_rows(values_only=True),
            start=1
        ):

            values = []

            for cell in row:
                if cell is not None:
                    values.append(str(cell).strip())

            # Skip completely empty rows
            if not values:
                continue

            text = " | ".join(values)

            rows.append({
                "sheet_name": sheet.title,
                "row_number": row_number,
                "text": text,
                "source_location": f"sheet {sheet.title}, row {row_number}"
            })

    workbook.close()

    return rows